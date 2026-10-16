"""Uncached quarantine workbook with explicit unavailable valuations."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.domains.pos.quarantine_report_schemas import QuarantineReportOverview
from app.domains.pos.xlsx_safe import xlsx_safe

UNIT_LABELS = {
    "tablet": "табл.",
    "capsule": "капс.",
    "ampoule": "амп.",
    "sachet": "пак.",
    "suppository": "супп.",
}
SUMMARY_LABELS = {"known": "Полная", "partial": "Неполная", "hidden": "Нет доступа"}
STATUS_LABELS = {"known": "Известна", "unknown": "Неизвестна", "hidden": "Нет доступа"}


def render_quarantine_report_xlsx(data: QuarantineReportOverview) -> bytes:
    workbook = Workbook()
    sheet: Worksheet = workbook.active
    sheet.title = "Карантин на дату"
    sheet.append(["Карантин на дату", data.on_date.isoformat(), "TJS"])
    sheet.append(
        ["Возвраты, ожидавшие решения на дату; не прибыль и не стоимость доступного остатка"]
    )
    sheet.append(["Полнота оценки", SUMMARY_LABELS[data.cost_status]])
    sheet.append(["Стоимость карантина", data.total_value])
    sheet.append(["Доступная известная стоимость", data.known_total_value])
    sheet.cell(4, 2).number_format = "0.00"
    sheet.cell(5, 2).number_format = "0.00"
    sheet.append(["Обычные упаковки", data.total_qty])
    for quantity in data.physical_quantities:
        sheet.append([UNIT_LABELS[quantity.unit], str(quantity.quantity)])
    sheet.append([])
    sheet.append(["Препарат", "Аптека", "Количество", "Единица", "Стоимость, TJS", "Оценка"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in data.rows:
        physical = row.physical_quantity
        sheet.append(
            [
                xlsx_safe(row.name),
                xlsx_safe(row.branch_name),
                str(physical.quantity) if physical else str(row.qty),
                UNIT_LABELS[physical.unit] if physical else "уп.",
                row.value,
                STATUS_LABELS[row.cost_status],
            ]
        )
        sheet.cell(sheet.max_row, 5).number_format = "0.00"
    sheet.freeze_panes = "A2"
    for column in ("A", "B"):
        sheet.column_dimensions[column].width = 30
    for column in ("C", "D", "E", "F"):
        sheet.column_dimensions[column].width = 22
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
