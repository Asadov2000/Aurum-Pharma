"""Uncached gross-profit workbook; partial sums and warehouse losses remain explicit."""

from datetime import UTC
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.domains.pos.profit_report_schemas import ProfitReportOverview
from app.domains.pos.xlsx_safe import xlsx_safe

COVERAGE = {"known": "Полная", "partial": "Неполная", "hidden": "Нет доступа"}
COST_LABELS = {"known": "Известна", "unknown": "Неизвестна", "hidden": "Нет доступа"}
EVENT_LABELS = {
    "sale": "Продажа",
    "refund": "Возврат денег",
    "zero_return": "Возврат без выплаты",
    "writeoff": "Складское списание",
}


def render_profit_report_xlsx(data: ProfitReportOverview) -> bytes:
    workbook = Workbook()
    sheet: Worksheet = workbook.active
    sheet.title = "Валовая прибыль"
    sheet.append(["Валовая прибыль продаж", str(data.date_from), str(data.date_to), "TJS"])
    sheet.append(["Не чистая прибыль. Складские списания показаны отдельно."])
    sheet.append(
        ["Известная часть относится только к строкам с доступной исторической стоимостью."]
    )
    sheet.append(["Полнота оценки продаж", COVERAGE[data.cost_status]])
    for label, value in (
        ("Выручка продаж после скидок", data.sales_revenue),
        ("Возвраты денег", data.refunds),
        ("Выручка за вычетом возвратов", data.net_revenue),
        ("Себестоимость за вычетом возвратов", data.net_sales_cost),
        ("Известная часть себестоимости", data.known_net_sales_cost),
        ("Валовая прибыль", data.gross_profit),
        ("Известная часть валовой прибыли", data.known_gross_profit),
        ("Складские списания", data.writeoff_cost),
        ("Известная стоимость списаний", data.known_writeoff_cost),
    ):
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 2).number_format = "0.00"
    sheet.append(["Полнота оценки списаний", COVERAGE[data.writeoff_cost_status]])
    sheet.append([])
    sheet.append(
        ["Дата UTC", "Операция", "Препарат", "Аптека", "Выручка", "Стоимость", "Прибыль", "Оценка"]
    )
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for row in data.rows:
        sheet.append(
            [
                row.occurred_at.astimezone(UTC).isoformat(),
                EVENT_LABELS[row.event_kind],
                xlsx_safe(row.name),
                xlsx_safe(row.branch_name),
                row.revenue,
                row.cost,
                row.gross_profit,
                COST_LABELS[row.cost_status],
            ]
        )
        for column in (5, 6, 7):
            sheet.cell(sheet.max_row, column).number_format = "0.00"
    sheet.freeze_panes = "A2"
    for column in ("A", "B", "C", "D"):
        sheet.column_dimensions[column].width = 30
    for column in ("E", "F", "G", "H"):
        sheet.column_dimensions[column].width = 22
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
