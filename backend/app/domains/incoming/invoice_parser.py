"""Bounded, read-only parsing of supplier spreadsheets into untrusted text cells."""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.errors import BusinessRuleError

MAX_INVOICE_FILE_BYTES = 10 * 1024 * 1024
MAX_INVOICE_ROWS = 500
MAX_INVOICE_COLUMNS = 32
MAX_INVOICE_CELL_LENGTH = 255
MAX_XLSX_ARCHIVE_MEMBERS = 2_048
MAX_XLSX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
INVOICE_FIELDS = (
    "source_name",
    "batch_number",
    "manufactured_at",
    "expires_at",
    "qty",
    "purchase_price",
    "sale_price",
)

_ALIASES: dict[str, tuple[str, ...]] = {
    "source_name": (
        "наименование",
        "номи ассортимент",
        "номи маҳсулот",
        "товар",
        "препарат",
        "наименование товара",
        "name",
        "product",
    ),
    "batch_number": ("серия", "партия", "номер партии", "batch", "lot"),
    "manufactured_at": ("дата производства", "дата изготовления", "manufactured_at"),
    "expires_at": ("срок годности", "годен до", "срок", "фаъолият", "expires_at", "expiry"),
    "qty": ("количество", "миқдор", "микдор", "кол-во", "кол во", "qty", "quantity"),
    "purchase_price": ("закупочная цена", "цена закупки", "нарх", "цена", "purchase_price"),
    "sale_price": ("цена продажи", "розничная цена", "sale_price"),
}


@dataclass(frozen=True, slots=True)
class InvoiceTable:
    headers: list[str]
    rows: list[list[str]]


def _invalid(reason: str, **details: int | list[int]) -> BusinessRuleError:
    return BusinessRuleError("Cannot read supplier invoice", details={"reason": reason, **details})


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date | datetime):
        result = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        result = str(int(value))
    else:
        result = str(value).strip()
    if len(result) > MAX_INVOICE_CELL_LENGTH:
        raise _invalid("invoice_file_invalid")
    return result


def _build_table(raw_rows: Sequence[Sequence[object]]) -> InvoiceTable:
    if not raw_rows:
        raise _invalid("invoice_file_invalid")
    if len(raw_rows[0]) > MAX_INVOICE_COLUMNS:
        raise _invalid("invoice_file_invalid")
    headers = [_cell(value) for value in raw_rows[0]]
    if not any(headers):
        raise _invalid("invoice_file_invalid")
    rows: list[list[str]] = []
    for values in raw_rows[1:]:
        if len(values) > MAX_INVOICE_COLUMNS:
            raise _invalid("invoice_file_invalid")
        row = [_cell(value) for value in values]
        if any(row):
            rows.append(row + [""] * (len(headers) - len(row)))
        if len(rows) > MAX_INVOICE_ROWS:
            raise _invalid("invoice_row_limit", limit=MAX_INVOICE_ROWS)
    if not rows:
        raise _invalid("invoice_file_invalid")
    return InvoiceTable(headers=headers, rows=rows)


def _parse_csv(raw: bytes) -> InvoiceTable:
    try:
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = raw.decode("cp1251")
        sample = decoded[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter, strict=True)
        rows: list[Sequence[object]] = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_INVOICE_ROWS + 2:
                raise _invalid("invoice_row_limit", limit=MAX_INVOICE_ROWS)
        return _build_table(rows)
    except (UnicodeDecodeError, csv.Error) as exc:
        raise _invalid("invoice_file_invalid") from exc


def _parse_xlsx(raw: bytes) -> InvoiceTable:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_ARCHIVE_MEMBERS:
                raise _invalid("invoice_file_invalid")
            if sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise _invalid("invoice_file_invalid")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            if not workbook.worksheets:
                raise _invalid("invoice_file_invalid")
            rows: list[list[object]] = []
            for row in workbook.worksheets[0].iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_INVOICE_ROWS + 2:
                    raise _invalid("invoice_row_limit", limit=MAX_INVOICE_ROWS)
            return _build_table(rows)
        finally:
            workbook.close()
    except BusinessRuleError:
        raise
    except Exception:
        # openpyxl and XML parsers can raise several different exceptions for
        # malformed supplier files. None of their messages belongs in an API.
        raise _invalid("invoice_file_invalid") from None


def parse_invoice(raw: bytes, filename: str | None) -> InvoiceTable:
    if not raw or len(raw) > MAX_INVOICE_FILE_BYTES:
        raise _invalid("invoice_file_invalid")
    extension = Path(filename or "").suffix.lower()
    if extension == ".csv":
        return _parse_csv(raw)
    if extension == ".xlsx":
        return _parse_xlsx(raw)
    raise _invalid("invoice_file_invalid")


def suggested_mapping(headers: list[str]) -> dict[str, int | None]:
    normalized = [" ".join(value.casefold().replace("_", " ").split()) for value in headers]
    mapping: dict[str, int | None] = {}
    used: set[int] = set()
    for field in INVOICE_FIELDS:
        match = next(
            (
                index
                for index, value in enumerate(normalized)
                if index not in used and value in _ALIASES[field]
            ),
            None,
        )
        mapping[field] = match
        if match is not None:
            used.add(match)
    return mapping


def validate_mapping(mapping: object, column_count: int) -> dict[str, int | None]:
    if not isinstance(mapping, dict) or set(mapping) != set(INVOICE_FIELDS):
        raise _invalid("invoice_mapping_invalid")
    result: dict[str, int | None] = {}
    used: set[int] = set()
    for field in INVOICE_FIELDS:
        index = mapping[field]
        if index is not None and (
            type(index) is not int or index < 0 or index >= column_count or index in used
        ):
            raise _invalid("invoice_mapping_invalid")
        result[field] = index
        if index is not None:
            used.add(index)
    if result["source_name"] is None:
        raise _invalid("invoice_mapping_invalid")
    return result
