"""Sales-summary aggregates, range filtering, the inverted-range guard, and the
empty-period file."""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import BusinessRuleError
from app.domains.pos.repository import POSRepository
from app.domains.pos.sales_summary_xlsx import render_sales_summary_xlsx
from app.domains.pos.service import POSService
from app.domains.pos.z_report_xlsx import render_z_report_xlsx
from tests.platform_access_helpers import create_test_platform_user


async def test_discounted_reports_reconcile_with_payments_and_refunds(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold(sale_price=10, batch_qty=10)
    service = POSService(POSRepository(db_session))
    shift = await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    sale = await service.create_sale(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        cashier_user_id=s["cashier"].id,
    )
    items, _ = await service.add_item(sale_id=sale.id, catalog_id=s["item"].id, qty=Decimal("3"))
    await service.repo.update_item(
        items[0], total_price=Decimal("25.00"), discount_amount=Decimal("5.00")
    )
    await service.repo.update_sale(sale, total_amount=Decimal("25.00"))
    await service.add_payment(sale_id=sale.id, payment_method="cash", amount=Decimal("25"))
    await service.complete(sale_id=sale.id)
    today = await _today(db_session)
    filters = {
        "tenant_id": s["tenant"].id,
        "date_from": today,
        "date_to": today,
        "branch_id": None,
    }
    reader = await create_test_platform_user(db_session, access_kind="developer")
    await db_session.execute(
        text(
            "SELECT set_config('app.user_id',:actor,true),"
            "set_config('app.tenant_id',:tenant,true)"
        ),
        {"actor": str(reader.id), "tenant": str(s["tenant"].id)},
    )
    summary = await service.build_sales_summary(**filters)
    overview = await service.get_sales_summary_overview(**filters)
    top = await service.get_top_products_overview(**filters, sort_by="revenue", limit=20)
    z = await service.build_z_report(shift.id)
    assert summary.gross_sales == overview.gross_sales == z.total_sales == Decimal("30")
    assert summary.net == overview.net == overview.daily[0].net == Decimal("25")
    assert overview.average_sale == top.rows[0].revenue == Decimal("25")
    assert summary.rows[0].gross == Decimal("30")
    assert summary.rows[0].net == summary.payment_breakdown.cash == Decimal("25")
    await service.refund(
        parent_sale_id=sale.id,
        items=[(items[0].id, Decimal("3"))],
        reason="other",
        comment=None,
        cashier_user_id=s["cashier"].id,
    )
    summary = await service.build_sales_summary(**filters)
    overview = await service.get_sales_summary_overview(**filters)
    assert summary.net == overview.net == overview.daily[0].net == Decimal("0")
    assert summary.total_refunds == Decimal("25")
    returned = next(row for row in summary.rows if row.kind == "return")
    assert returned.net == Decimal("25")
    z = await service.build_z_report(shift.id)
    assert z.total_sales - z.total_discounts - z.total_refunds == Decimal("0")
    # Historical all-branch exports retain provenance after branch archival.
    s["branch"].is_active = False
    await db_session.flush()
    summary = await service.build_sales_summary(**filters)
    assert summary.show_branch_column is True
    assert all(row.branch_name == s["branch"].name for row in summary.rows)


async def test_report_exports_render_tenant_local_timestamps(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold()
    instant = datetime.now(UTC).replace(hour=20, minute=15, second=0, microsecond=0)
    service = POSService(POSRepository(db_session), now=lambda: instant)
    shift = await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    await _complete_sale(service, s, qty=Decimal("1"), payments=[("cash", Decimal("10"))])
    local_date = instant.date() + timedelta(days=1)
    assert await service.report_local_date(s["tenant"].id) == local_date
    summary = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=local_date, date_to=local_date, branch_id=None
    )
    workbook = load_workbook(BytesIO(render_sales_summary_xlsx(summary)))
    assert workbook["Детализация"].cell(2, 1).value == f"{local_date:%d.%m.%Y} 01:15"
    z = await service.build_z_report(shift.id)
    workbook = load_workbook(BytesIO(render_z_report_xlsx(z)))
    values = dict(workbook.active.iter_rows(min_col=1, max_col=2, values_only=True))
    assert values["Открыта"] == z.opened_at.strftime("%d.%m.%Y %H:%M")
    assert z.opened_at.utcoffset() == timedelta(hours=5)


async def _today(db: AsyncSession) -> date:
    # Reports interpret dates in the tenant tz (Asia/Dushanbe by default), so
    # "today" here must be the local date — not UTC, which disagrees by a day
    # between 19:00–23:59 UTC.
    return (
        await db.execute(text("SELECT (now() AT TIME ZONE 'Asia/Dushanbe')::date"))
    ).scalar_one()


async def _complete_sale(  # type: ignore[no-untyped-def]
    service: POSService,
    s,
    *,
    qty,
    payments,
    completed_at: datetime | None = None,
    is_test: bool = False,
):
    sale = await service.create_sale(
        tenant_id=s["tenant"].id, register_id=s["register"].id, cashier_user_id=s["cashier"].id
    )
    created, _ = await service.add_item(sale_id=sale.id, catalog_id=s["item"].id, qty=qty)
    for method, amount in payments:
        if method == "cash":
            await service.add_payment(
                sale_id=sale.id,
                payment_method=method,
                amount=amount,
            )
        else:
            await service.repo.insert_payment(
                tenant_id=s["tenant"].id,
                sale_id=sale.id,
                payment_method=method,
                amount=amount,
            )
    if is_test:
        await service.repo.update_sale(sale, is_test=True)
    completion_service = (
        service
        if completed_at is None
        else POSService(service.repo, now=lambda value=completed_at: value)
    )
    await completion_service.complete(sale_id=sale.id)
    return sale, created


async def _complete_historical_bank_transfer(
    service: POSService,
    s,
):  # type: ignore[no-untyped-def]
    sale = await service.create_sale(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        cashier_user_id=s["cashier"].id,
    )
    await service.add_item(
        sale_id=sale.id,
        catalog_id=s["item"].id,
        qty=Decimal("1"),
    )
    await service.repo.insert_payment(
        tenant_id=s["tenant"].id,
        sale_id=sale.id,
        payment_method="bank_transfer",
        amount=Decimal("10"),
    )
    await service.complete(sale_id=sale.id)


async def test_sales_summary_aggregates_and_breakdown(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold(sale_price=Decimal("10"), batch_qty=100)
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )

    sale_a, created_a = await _complete_sale(
        service, s, qty=Decimal("2"), payments=[("cash", Decimal("20"))]
    )
    await _complete_sale(service, s, qty=Decimal("1"), payments=[("card", Decimal("10"))])
    await _complete_sale(service, s, qty=Decimal("1"), payments=[("qr", Decimal("10"))])
    await _complete_sale(
        service, s, qty=Decimal("1"), payments=[("cash", Decimal("5")), ("card", Decimal("5"))]
    )
    await _complete_historical_bank_transfer(service, s)
    await service.refund(
        parent_sale_id=sale_a.id,
        items=[(created_a[0].id, Decimal("1"))],
        reason="quality_issue",
        comment=None,
        cashier_user_id=s["cashier"].id,
    )

    today = await _today(db_session)
    data = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )

    assert data.gross_sales == Decimal("60.00")
    assert data.total_discounts == Decimal("0")
    assert data.total_refunds == Decimal("10.00")
    assert data.net == Decimal("50.00")
    assert data.sales_count == 5
    assert data.returns_count == 1

    assert data.payment_breakdown.cash == Decimal("20.00")
    assert data.payment_breakdown.card == Decimal("10.00")
    assert data.payment_breakdown.qr == Decimal("10.00")
    assert data.payment_breakdown.mixed == Decimal("10.00")
    assert data.payment_breakdown.bank_transfer == Decimal("10.00")

    # Five sales + one return row, returns shown separately (not folded into sales).
    assert len(data.rows) == 6
    kinds = sorted(r.kind for r in data.rows)
    assert kinds == ["return", "sale", "sale", "sale", "sale", "sale"]

    overview = await service.get_sales_summary_overview(
        tenant_id=s["tenant"].id,
        date_from=today,
        date_to=today,
        branch_id=None,
    )
    assert overview.gross_sales == Decimal("60.00")
    assert overview.net == Decimal("50.00")
    assert overview.average_sale == Decimal("12.00")
    assert overview.payment_breakdown.mixed == Decimal("10.00")
    assert len(overview.daily) == 1
    assert overview.daily[0].day == today
    assert overview.daily[0].gross_sales == Decimal("60.00")
    assert overview.daily[0].total_refunds == Decimal("10.00")
    assert overview.daily[0].net == Decimal("50.00")
    assert overview.daily[0].sales_count == 5
    assert overview.daily[0].returns_count == 1


async def test_sales_summary_range_excludes_sales_outside_period(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold(sale_price=Decimal("10"))
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    await _complete_sale(service, s, qty=Decimal("1"), payments=[("cash", Decimal("10"))])

    today = await _today(db_session)
    # A window entirely before today must not see today's sale.
    past = await service.build_sales_summary(
        tenant_id=s["tenant"].id,
        date_from=today - timedelta(days=10),
        date_to=today - timedelta(days=1),
        branch_id=None,
    )
    assert past.gross_sales == Decimal("0")
    assert past.sales_count == 0
    assert past.rows == []

    # Today's window sees it.
    now = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )
    assert now.gross_sales == Decimal("10.00")
    assert now.sales_count == 1


async def test_sales_summary_full_refund_nets_to_zero(
    db_session: AsyncSession, pos_scaffold
) -> None:
    """A 100 sale fully refunded in the same period: gross=100, refunds=100,
    net=0 — not gross=0/net=-100. The voided parent still counts as a sale."""
    s = await pos_scaffold(sale_price=Decimal("100"), batch_qty=10)
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    sale, created = await _complete_sale(
        service, s, qty=Decimal("1"), payments=[("cash", Decimal("100"))]
    )
    # Full refund (the only unit) → parent voided + a return document for 100.
    await service.refund(
        parent_sale_id=sale.id,
        items=[(created[0].id, Decimal("1"))],
        reason="quality_issue",
        comment=None,
        cashier_user_id=s["cashier"].id,
    )

    today = await _today(db_session)
    data = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )

    assert data.gross_sales == Decimal("100.00")
    assert data.total_refunds == Decimal("100.00")
    assert data.net == Decimal("0")  # money in then back out
    assert data.sales_count == 1
    assert data.returns_count == 1
    assert data.payment_breakdown.cash == Decimal("100.00")

    # Both rows are visible: the cancelled sale and the refund.
    kinds = sorted(r.kind for r in data.rows)
    assert kinds == ["return", "voided"]


async def test_sales_summary_rejects_inverted_range(db_session: AsyncSession, pos_scaffold) -> None:
    s = await pos_scaffold()
    service = POSService(POSRepository(db_session))
    today = await _today(db_session)
    with pytest.raises(BusinessRuleError):
        await service.get_sales_summary_xlsx(
            tenant_id=s["tenant"].id,
            date_from=today,
            date_to=today - timedelta(days=1),
            branch_id=None,
        )
    with pytest.raises(BusinessRuleError):
        await service.get_sales_summary_overview(
            tenant_id=s["tenant"].id,
            date_from=today,
            date_to=today - timedelta(days=1),
            branch_id=None,
        )
    with pytest.raises(BusinessRuleError):
        await service.get_sales_summary_overview(
            tenant_id=s["tenant"].id,
            date_from=today - timedelta(days=366),
            date_to=today,
            branch_id=None,
        )
    with pytest.raises(BusinessRuleError):
        await service.get_sales_summary_xlsx(
            tenant_id=s["tenant"].id,
            date_from=today - timedelta(days=366),
            date_to=today,
            branch_id=None,
        )


async def test_sales_summary_empty_period_is_valid_zero_file(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold()
    service = POSService(POSRepository(db_session))
    today = await _today(db_session)

    data = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )
    assert data.gross_sales == Decimal("0")
    assert data.sales_count == 0
    assert data.rows == []

    overview = await service.get_sales_summary_overview(
        tenant_id=s["tenant"].id,
        date_from=today,
        date_to=today,
        branch_id=None,
    )
    assert overview.net == Decimal("0")
    assert overview.average_sale == Decimal("0.00")
    assert overview.daily == []

    xlsx = await service.get_sales_summary_xlsx(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )
    assert xlsx[:2] == b"PK"  # valid workbook, not an error
    assert len(xlsx) > 500


async def test_sales_summary_dates_use_tenant_timezone(
    db_session: AsyncSession, pos_scaffold
) -> None:
    """A 01:00 Dushanbe sale (= 20:00 UTC the previous day) belongs to the local
    day, not the previous UTC day. Default tz is Asia/Dushanbe."""
    s = await pos_scaffold(sale_price=Decimal("10"))
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    # 2026-06-15 01:00 Dushanbe == 2026-06-14 20:00 UTC.
    await _complete_sale(
        service,
        s,
        qty=Decimal("1"),
        payments=[("cash", Decimal("10"))],
        completed_at=datetime(2026, 6, 14, 20, 0, tzinfo=UTC),
    )

    local_day = await service.build_sales_summary(
        tenant_id=s["tenant"].id,
        date_from=date(2026, 6, 15),
        date_to=date(2026, 6, 15),
        branch_id=None,
    )
    assert local_day.gross_sales == Decimal("10.00")
    assert local_day.sales_count == 1

    local_overview = await service.get_sales_summary_overview(
        tenant_id=s["tenant"].id,
        date_from=date(2026, 6, 15),
        date_to=date(2026, 6, 15),
        branch_id=None,
    )
    assert [item.day for item in local_overview.daily] == [date(2026, 6, 15)]

    prev_utc_day = await service.build_sales_summary(
        tenant_id=s["tenant"].id,
        date_from=date(2026, 6, 14),
        date_to=date(2026, 6, 14),
        branch_id=None,
    )
    assert prev_utc_day.gross_sales == Decimal("0")
    assert prev_utc_day.sales_count == 0


async def test_sales_summary_excludes_test_sales(db_session: AsyncSession, pos_scaffold) -> None:
    """is_test sales are not real money — excluded from totals and the receipt
    list; the normal sale is counted."""
    s = await pos_scaffold(sale_price=Decimal("10"))
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    await _complete_sale(service, s, qty=Decimal("1"), payments=[("cash", Decimal("10"))])
    await _complete_sale(
        service,
        s,
        qty=Decimal("1"),
        payments=[("cash", Decimal("10"))],
        is_test=True,
    )

    today = await _today(db_session)
    data = await service.build_sales_summary(
        tenant_id=s["tenant"].id, date_from=today, date_to=today, branch_id=None
    )
    assert data.gross_sales == Decimal("10.00")  # only the real sale
    assert data.sales_count == 1
    assert len(data.rows) == 1

    rows, total = await service.list_sales(
        tenant_id=s["tenant"].id,
        cashier_id=None,
        branch_id=None,
        register_id=None,
        receipt_number=None,
        date_from=None,
        date_to=None,
        has_refund=None,
        min_total=None,
        max_total=None,
        page=1,
        page_size=50,
    )
    assert total == 1  # the test sale is hidden from receipt search
