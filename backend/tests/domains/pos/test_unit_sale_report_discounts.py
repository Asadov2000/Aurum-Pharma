"""Financial reports read immutable unit-sale receipts through the real app role."""

import json
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID, uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from app.core.deps import CurrentUser
from app.core.errors import BusinessRuleError, NotFoundError
from app.core.storage import ensure_bucket, get_object, put_object
from app.domains.pos.mixed_cash_refund_schemas import MixedCashRefundResultV4
from app.domains.pos.mixed_documents import MixedCheckoutResultV4
from app.domains.pos.mixed_return_allocation import LegacyReturnSelection
from app.domains.pos.physical_cash_refund_schemas import PhysicalCashRefundResultV3
from app.domains.pos.physical_documents import PhysicalCheckoutResultV3
from app.domains.pos.physical_return_allocation import PhysicalReturnSelection
from app.domains.pos.repository import _SALE_DISCOUNT_SQL, POSRepository
from app.domains.pos.sales_summary_xlsx import render_sales_summary_xlsx
from app.domains.pos.schemas import SalesSummaryData
from app.domains.pos.service import POSService
from app.domains.pos.z_report_xlsx import render_z_report_xlsx
from tests.domains.incoming.test_unit_acceptance import db_session as receiving_session_fixture
from tests.domains.incoming.test_unit_acceptance import seed_receiving_actor, set_receiving_context
from tests.domains.inventory.test_unit_ledger import seed_ledger_context
from tests.domains.pos.test_mixed_cash_checkout import complete as complete_mixed
from tests.domains.pos.test_mixed_cash_refund import facts as mixed_facts
from tests.domains.pos.test_mixed_cash_refund import refund as refund_mixed
from tests.domains.pos.test_mixed_physical_checkout_preparation import legacy_with_stock
from tests.domains.pos.test_mixed_return_journal import MixedSource, mixed_source
from tests.domains.pos.test_physical_cart_api import cart_actor
from tests.domains.pos.test_physical_cash_checkout import complete
from tests.domains.pos.test_physical_cash_refund import facts as physical_facts
from tests.domains.pos.test_physical_cash_refund import refund as refund_physical
from tests.domains.pos.test_unit_physical_cart import put_physical
from tests.domains.pos.test_unit_physical_checkout_preparation import checkout_source, saved_request
from tests.domains.pos.test_unit_return import runtime

db_session = receiving_session_fixture


async def report_day(session: AsyncSession) -> date:
    result = await session.scalar(text("SELECT (now() AT TIME ZONE 'Asia/Dushanbe')::date"))
    assert isinstance(result, date)
    return result


def checkout_gross(result: MixedCheckoutResultV4 | PhysicalCheckoutResultV3) -> Decimal:
    # Fixtures use different pack prices. Reconcile against the posted retail
    # amounts, independently of the report's net-plus-discount calculation.
    return sum((line.gross_amount for line in result.lines), Decimal("0.00"))


async def reconciled_reports(
    session: AsyncSession,
    shift_id: UUID,
    *,
    gross: Decimal,
    discount: Decimal,
    refunds: Decimal,
) -> SalesSummaryData:
    context = runtime(session)
    day = await report_day(session)
    service = POSService(POSRepository(session))
    summary = await service.build_sales_summary(
        tenant_id=context.stock.tenant_id,
        date_from=day,
        date_to=day,
        branch_id=context.stock.branch_id,
    )
    overview = await service.get_sales_summary_overview(
        tenant_id=context.stock.tenant_id,
        date_from=day,
        date_to=day,
        branch_id=context.stock.branch_id,
    )
    z = await service.build_z_report(shift_id)
    assert summary.gross_sales == overview.gross_sales == z.total_sales == gross
    assert summary.total_discounts == overview.total_discounts == z.total_discounts == discount
    assert overview.daily[0].gross_sales == gross
    assert overview.daily[0].total_discounts == discount
    assert summary.total_refunds == overview.total_refunds == z.total_refunds == refunds
    assert summary.net == overview.net == overview.daily[0].net == gross - discount - refunds
    assert z.total_sales - z.total_discounts - z.total_refunds == summary.net
    z_workbook = load_workbook(BytesIO(render_z_report_xlsx(z)))
    worksheet = z_workbook.active
    assert worksheet is not None
    values = dict(worksheet.iter_rows(min_col=1, max_col=2, values_only=True))
    assert Decimal(str(values["Скидки"])) == discount
    assert Decimal(str(values["Валовая сумма"])) == gross
    return summary


@asynccontextmanager
async def report_actor(session: AsyncSession, actor_id: UUID) -> AsyncIterator[CurrentUser]:
    context = runtime(session)
    await set_receiving_context(session, context.stock, actor_id)
    async with cart_actor(session) as actor:
        actor.user_id = actor_id
        actor.permissions = {"reports.view", "reports.export"}
        actor.permission_scopes = {
            code: frozenset({context.stock.branch_id}) for code in actor.permissions
        }
        try:
            yield actor
        finally:
            await set_receiving_context(session, context.stock, context.actor_id)


@pytest.mark.parametrize("kind", ["physical", "mixed"])
async def test_unit_sale_and_partial_then_full_refund_reconcile_every_report(
    db_session: AsyncSession, kind: Literal["physical", "mixed"]
) -> None:
    first: MixedCashRefundResultV4 | PhysicalCashRefundResultV3
    second: MixedCashRefundResultV4 | PhysicalCashRefundResultV3
    original: MixedCheckoutResultV4 | PhysicalCheckoutResultV3
    if kind == "mixed":
        source = await mixed_source(db_session, fractional_legacy=True)
        original = source.original
        discount = Decimal("0.08")
        first = await refund_mixed(
            db_session,
            source,
            (
                LegacyReturnSelection(source.legacy_id, Decimal("0.001")),
                PhysicalReturnSelection(source.physical_id, "unit", 1),
            ),
        )
        partial_discount = sum(
            (line.discount_amount for line in first.receipt.lines), Decimal("0.00")
        )
        original_id, shift_id = original.sale_id, original.shift_id
    else:
        physical = await checkout_source(db_session)
        item = await put_physical(
            db_session,
            physical.sale.id,
            sale_format="pack",
            quantity=2,
            discount=Decimal("3.01"),
        )
        original = await complete(db_session, await saved_request(db_session, physical.sale))
        discount = Decimal("3.01")
        original_id, shift_id = original.sale_id, original.shift_id
        first = await refund_physical(db_session, original_id, item.id, shift_id)
        partial_discount = first.receipt.lines[0].discount_amount

    gross = checkout_gross(original)
    await db_session.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    # Validate the first operation, then restore the normal transaction mode:
    # the next journal inserts its parent before its lines and slices.
    await db_session.execute(text("SET CONSTRAINTS ALL DEFERRED"))
    original_row = await db_session.scalar(
        text("SELECT to_jsonb(sale) FROM sale WHERE id=:id"), {"id": original_id}
    )
    summary = await reconciled_reports(
        db_session, shift_id, gross=gross, discount=discount, refunds=first.total_amount
    )
    returned = next(row for row in summary.rows if row.kind == "return")
    assert returned.discount == partial_discount
    assert returned.gross == first.total_amount + partial_discount
    assert returned.net == first.total_amount

    if kind == "mixed":
        second = await refund_mixed(
            db_session,
            source,
            (
                LegacyReturnSelection(source.legacy_id, Decimal("0.002")),
                PhysicalReturnSelection(source.physical_id, "unit", 59),
            ),
        )
    else:
        second = await refund_physical(db_session, original_id, item.id, shift_id, quantity=59)
    await db_session.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    assert first.total_amount + second.total_amount == gross - discount
    before = await (mixed_facts(db_session) if kind == "mixed" else physical_facts(db_session))
    summary = await reconciled_reports(
        db_session, shift_id, gross=gross, discount=discount, refunds=gross - discount
    )
    assert summary.net == Decimal("0.00")
    assert summary.sales_count == 1 and summary.returns_count == 2
    assert sum(row.discount for row in summary.rows if row.kind == "return") == discount
    workbook = load_workbook(BytesIO(render_sales_summary_xlsx(summary)))
    detail = list(workbook["Детализация"].iter_rows(min_row=2, values_only=True))
    assert len(detail) == 3
    for exported, row in zip(detail, summary.rows, strict=True):
        assert exported[2] == row.cashier_name
        assert Decimal(str(exported[4])) == row.gross
        assert Decimal(str(exported[5])) == row.discount
        assert Decimal(str(exported[6])) == row.net
    assert (
        await db_session.scalar(
            text("SELECT to_jsonb(sale) FROM sale WHERE id=:id"), {"id": original_id}
        )
        == original_row
    )
    after = await (mixed_facts(db_session) if kind == "mixed" else physical_facts(db_session))
    assert after == before


async def test_report_only_actor_reads_mixed_discounts_without_private_journal_access(
    client: AsyncClient,
    db_session: AsyncSession,
    db_engine: AsyncEngine,
    maintenance_engine: AsyncEngine,
) -> None:
    context = runtime(db_session)
    reader_id = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=("reports.view", "reports.export")
    )
    source = await mixed_source(db_session, fractional_legacy=True)
    await db_session.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    before = await mixed_facts(db_session)
    day = await report_day(db_session)
    params = {"from": str(day), "to": str(day), "branch_id": str(context.stock.branch_id)}
    async with report_actor(db_session, reader_id) as actor:
        for table in ("pos_mixed_checkout", "unit_stock_cart_command_v2"):
            assert (
                await db_session.scalar(
                    text(f"SELECT count(*) FROM {table} WHERE tenant_id=:tenant"),
                    {"tenant": context.stock.tenant_id},
                )
                == 0
            )
        summary = await client.get("/api/v1/reports/sales-summary", params=params)
        assert summary.status_code == 200
        assert Decimal(summary.json()["total_discounts"]) == Decimal("0.08")
        assert Decimal(summary.json()["gross_sales"]) == checkout_gross(source.original)
        z = await client.get(f"/api/v1/shifts/{source.original.shift_id}/z-report")
        assert z.status_code == 200
        assert Decimal(z.json()["totals"]["discounts_total"]) == Decimal("0.08")
        export = await client.get("/api/v1/reports/sales-summary.xlsx", params=params)
        assert export.status_code == 200 and export.content[:2] == b"PK"
        actor.permission_scopes = {code: frozenset({uuid4()}) for code in actor.permissions}
        assert (await client.get("/api/v1/reports/sales-summary", params=params)).status_code == 403
        assert (
            await client.get(f"/api/v1/shifts/{source.original.shift_id}/z-report")
        ).status_code == 403
    assert await mixed_facts(db_session) == before


async def test_foreign_tenant_cannot_read_unit_sale_report_discounts(
    db_session: AsyncSession, db_engine: AsyncEngine, maintenance_engine: AsyncEngine
) -> None:
    source = await mixed_source(db_session, fractional_legacy=True)
    context = runtime(db_session)
    day = await report_day(db_session)
    async with AsyncSession(db_engine) as support:
        async with support.begin():
            await support.execute(text("SELECT set_config('app.support_session','true',true)"))
            foreign = await seed_ledger_context(support)
    actor_id = await seed_receiving_actor(
        db_engine, maintenance_engine, foreign, permissions=("reports.view",)
    )
    await set_receiving_context(db_session, foreign, actor_id)
    try:
        service = POSService(POSRepository(db_session))
        summary = await service.build_sales_summary(
            tenant_id=context.stock.tenant_id,
            date_from=day,
            date_to=day,
            branch_id=context.stock.branch_id,
        )
        assert summary.rows == []
        assert summary.gross_sales == summary.total_discounts == Decimal("0.00")
        with pytest.raises(NotFoundError):
            await service.build_z_report(source.original.shift_id)
    finally:
        await set_receiving_context(db_session, context.stock, context.actor_id)


@pytest.mark.parametrize(
    "version,kind",
    [
        (4, "physical_receipt"),
        (3, "mixed_receipt"),
        (4, "physical_cash_refund_receipt"),
        (3, "mixed_cash_refund_receipt"),
        ("4", "mixed_receipt"),
        (4, "unrecognized_receipt"),
    ],
)
async def test_unknown_or_mismatched_snapshot_pair_preserves_legacy_source(
    db_session: AsyncSession, version: int | str, kind: str
) -> None:
    source = await mixed_source(db_session, fractional_legacy=True)
    # Project an unsupported snapshot in a read-only CTE. Committed receipts
    # remain untouched, and the fallback still reads the actual legacy SaleItem.
    result = await db_session.scalar(
        text(f"""
            WITH s AS (
                SELECT id, tenant_id, CAST(:payload AS jsonb) AS receipt_snapshot
                FROM sale WHERE id=:sale_id
            ) SELECT ({_SALE_DISCOUNT_SQL}) FROM s
        """),
        {
            "sale_id": source.original.sale_id,
            "payload": json.dumps(
                {
                    "schema_version": version,
                    "document_kind": kind,
                    "lines": [{"discount_amount": "999.00"}],
                }
            ),
        },
    )
    assert result == Decimal("0.07")


@pytest.mark.parametrize("lines", [None, [], {}, [{}], [{"discount_amount": None}]])
async def test_matching_snapshot_with_missing_money_fails_without_legacy_fallback(
    db_session: AsyncSession, lines: object
) -> None:
    source = await mixed_source(db_session, fractional_legacy=True)
    payload = json.dumps({"schema_version": 4, "document_kind": "mixed_receipt", "lines": lines})
    savepoint = await db_session.begin_nested()
    try:
        with pytest.raises(DBAPIError):
            await db_session.scalar(
                text(f"""
                    WITH s AS (
                        SELECT id, tenant_id, CAST(:payload AS jsonb) AS receipt_snapshot
                        FROM sale WHERE id=:sale_id
                    ) SELECT ({_SALE_DISCOUNT_SQL}) FROM s
                """),
                {"sale_id": source.original.sale_id, "payload": payload},
            )
    finally:
        await savepoint.rollback()


async def test_rejected_training_draft_has_no_reported_discounts_or_revenue(
    db_session: AsyncSession,
) -> None:
    source = await checkout_source(db_session)
    await put_physical(db_session, source.sale.id, quantity=2, discount=Decimal("0.01"))
    service = POSService(POSRepository(db_session))
    # Training cannot post physical stock or cash; completed history is never
    # rewritten to manufacture a reporting fixture.
    await service.repo.update_sale(source.sale, is_test=True)
    day = await report_day(db_session)
    context = runtime(db_session)
    filters = {
        "tenant_id": context.stock.tenant_id,
        "date_from": day,
        "date_to": day,
        "branch_id": context.stock.branch_id,
    }
    draft = await service.build_sales_summary(**filters)
    assert draft.rows == [] and draft.gross_sales == Decimal("0.00")
    request = await saved_request(db_session, source.sale)
    before = await physical_facts(db_session)
    with pytest.raises(BusinessRuleError, match="Physical test sales"):
        await complete(db_session, request)
    assert await physical_facts(db_session) == before
    await db_session.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    summary = await service.build_sales_summary(**filters)
    overview = await service.get_sales_summary_overview(**filters)
    z = await service.build_z_report(source.sale.shift_id)
    assert summary.rows == [] and overview.daily == []
    assert summary.total_discounts == overview.total_discounts == z.total_discounts == Decimal("0")
    assert summary.gross_sales == overview.gross_sales == z.total_sales == Decimal("0")
    assert summary.sales_count == z.sales_count == 0


async def test_closed_shift_export_ignores_previous_cache_and_reuses_new_workbook(
    db_session: AsyncSession,
) -> None:
    source = await mixed_source(db_session, fractional_legacy=True)
    context = runtime(db_session)
    service = POSService(POSRepository(db_session))
    await service.close_shift(
        shift_id=source.original.shift_id,
        closing_cash_actual=source.original.total_amount,
        closed_by_user_id=context.actor_id,
        allowed_branch_ids={context.stock.branch_id},
    )
    report = await service.build_z_report(source.original.shift_id)
    suffix = f"{report.report_timezone}/{report.shift_id}.xlsx"
    bucket = ensure_bucket()
    stale = b"previous cached report without physical discounts"
    put_object(
        object_name=f"zreports/v2/{suffix}",
        data=stale,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    first = await service.get_z_report_xlsx(source.original.shift_id)
    second = await service.get_z_report_xlsx(source.original.shift_id)
    assert first[:2] == b"PK" and first == second
    assert get_object(f"{bucket}/zreports/v3/{suffix}") == first
    assert get_object(f"{bucket}/zreports/v2/{suffix}") == stale
    workbook = load_workbook(BytesIO(first))
    worksheet = workbook.active
    assert worksheet is not None
    values = dict(worksheet.iter_rows(min_col=1, max_col=2, values_only=True))
    assert Decimal(str(values["Скидки"])) == Decimal("0.08")


@pytest.mark.parametrize("kind", ["physical", "mixed"])
async def test_next_day_refund_keeps_original_day_and_shift_discounts(
    db_session: AsyncSession, monkeypatch: pytest.MonkeyPatch, kind: Literal["physical", "mixed"]
) -> None:
    source = await checkout_source(db_session)
    context = runtime(db_session)
    service = POSService(POSRepository(db_session))
    previous = datetime.now(UTC) - timedelta(days=1)
    # Only editable creation facts are dated before checkout. Completion and its
    # immutable receipt are then written by the normal coordinator with this clock.
    await service.repo.update_sale(source.sale, created_at=previous - timedelta(minutes=1))
    old_shift = await service.repo.get_shift(source.sale.shift_id)
    assert old_shift is not None
    await service.repo.update_shift(old_shift, opened_at=previous - timedelta(minutes=2))
    if kind == "mixed":
        legacy = await legacy_with_stock(db_session, source.sale)
        await service.repo.update_item(
            legacy, discount_amount=Decimal("0.01"), total_price=Decimal("0.01")
        )
        await service.repo.update_sale(source.sale, total_amount=Decimal("0.01"))
    item = await put_physical(
        db_session,
        source.sale.id,
        position=2 if kind == "mixed" else 1,
        sale_format="pack",
        quantity=2,
        discount=Decimal("3.01"),
    )
    original: MixedCheckoutResultV4 | PhysicalCheckoutResultV3
    with monkeypatch.context() as clock:
        clock.setattr(f"app.domains.pos.{kind}_cash_checkout_service.utc_now", lambda: previous)
        request = await saved_request(db_session, source.sale)
        if kind == "mixed":
            original = await complete_mixed(db_session, request)
        else:
            original = await complete(db_session, request)
    await service.close_shift(
        shift_id=original.shift_id,
        closing_cash_actual=original.total_amount,
        closed_by_user_id=context.actor_id,
        allowed_branch_ids={context.stock.branch_id},
    )
    new_shift = await service.open_shift(
        tenant_id=context.stock.tenant_id,
        register_id=original.register_id,
        opened_by_user_id=context.actor_id,
        opening_cash=original.total_amount,
        allowed_branch_ids={context.stock.branch_id},
    )
    refund: MixedCashRefundResultV4 | PhysicalCashRefundResultV3
    if isinstance(original, MixedCheckoutResultV4):
        mixed = MixedSource(original, legacy.id, item.id, legacy.batch_id)
        refund = await refund_mixed(
            db_session, mixed, mixed.selections("0.001", 60), shift_id=new_shift.id
        )
        discount = Decimal("3.02")
    else:
        refund = await refund_physical(
            db_session, original.sale_id, item.id, new_shift.id, quantity=60
        )
        discount = Decimal("3.01")
    gross = checkout_gross(original)
    await db_session.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    before = await (mixed_facts(db_session) if kind == "mixed" else physical_facts(db_session))
    # Reports use the pharmacy's day; derive it with PostgreSQL, as the report does.
    old_day = await db_session.scalar(
        text("SELECT (CAST(:instant AS timestamptz) AT TIME ZONE 'Asia/Dushanbe')::date"),
        {"instant": previous},
    )
    assert isinstance(old_day, date)
    for day, expected_gross, expected_discount, expected_refunds in (
        (old_day, gross, discount, Decimal("0.00")),
        (await report_day(db_session), Decimal("0.00"), Decimal("0.00"), refund.total_amount),
    ):
        summary = await service.build_sales_summary(
            tenant_id=context.stock.tenant_id,
            date_from=day,
            date_to=day,
            branch_id=context.stock.branch_id,
        )
        overview = await service.get_sales_summary_overview(
            tenant_id=context.stock.tenant_id,
            date_from=day,
            date_to=day,
            branch_id=context.stock.branch_id,
        )
        assert summary.gross_sales == overview.gross_sales == expected_gross
        assert summary.total_discounts == overview.total_discounts == expected_discount
        assert summary.total_refunds == overview.total_refunds == expected_refunds
        assert (
            summary.net
            == overview.daily[0].net
            == (expected_gross - expected_discount - expected_refunds)
        )
    original_z = await service.build_z_report(original.shift_id)
    refund_z = await service.build_z_report(new_shift.id)
    assert original_z.total_sales == gross and original_z.total_discounts == discount
    assert original_z.total_refunds == Decimal("0.00")
    assert refund_z.total_sales == refund_z.total_discounts == Decimal("0.00")
    assert refund_z.total_refunds == original.total_amount == refund.total_amount
    after = await (mixed_facts(db_session) if kind == "mixed" else physical_facts(db_session))
    assert after == before
