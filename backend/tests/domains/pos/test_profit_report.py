"""Profit reads reconcile cash, immutable costs and independent warehouse losses."""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from app.domains.inventory.repository import InventoryRepository
from app.domains.inventory.service import InventoryService
from app.domains.pos.profit_report_repository import ProfitReportRepository
from app.domains.pos.profit_report_schemas import ProfitReportSnapshot
from tests.domains.customer_returns.test_mixed_queue import cash_source as mixed_cash
from tests.domains.customer_returns.test_physical_disposition_helpers import (
    confirmed,
    unchanged_facts,
)
from tests.domains.customer_returns.test_physical_queue import cash_source as physical_cash
from tests.domains.incoming.test_unit_acceptance import db_session as receiving_session_fixture
from tests.domains.incoming.test_unit_acceptance import seed_receiving_actor, set_receiving_context
from tests.domains.inventory.test_physical_writeoff_cancellation import cancel
from tests.domains.inventory.test_physical_writeoff_posting import post, selection
from tests.domains.inventory.test_unit_cost_ledger_v2 import priced_source
from tests.domains.pos.test_mixed_zero_return import intake
from tests.domains.pos.test_mixed_zero_return import source as mixed_zero
from tests.domains.pos.test_unit_return import runtime
from tests.domains.pos.test_unit_sale_report_discounts import (
    checkout_gross,
    reconciled_reports,
    report_actor,
    report_day,
)
from tests.domains.pos.test_unit_stock_on_date import legacy_batch

db_session = receiving_session_fixture
PATH = "/api/v1/reports/gross-profit"


async def snapshot(
    session: AsyncSession,
    *,
    start: datetime | None = None,
    end: datetime | None = None,
    tenant_id: UUID | None = None,
    branch_id: UUID | None = None,
    limit: int | None = None,
    offset: int = 0,
) -> ProfitReportSnapshot:
    now = await session.scalar(text("SELECT clock_timestamp()"))
    assert isinstance(now, datetime)
    context = runtime(session)
    return await ProfitReportRepository(session).snapshot(
        tenant_id=tenant_id or context.stock.tenant_id,
        start=start or now - timedelta(days=1),
        end=end or now + timedelta(days=1),
        branch_id=branch_id or context.stock.branch_id,
        limit=limit,
        offset=offset,
    )


@pytest.mark.parametrize("kind", ["physical", "mixed"])
@pytest.mark.parametrize("returned", [1, 31])
async def test_cash_refunds_reconcile_summary_z_and_saved_costs(
    db_session: AsyncSession, kind: str, returned: int
) -> None:
    original, refund = (
        await physical_cash(db_session, returned=returned)
        if kind == "physical"
        else await mixed_cash(db_session, returned=returned)
    )
    before = await unchanged_facts(db_session)
    result = await snapshot(db_session)
    gross = checkout_gross(original)
    discount = sum((row.discount_amount for row in original.lines), Decimal("0"))
    summary = await reconciled_reports(
        db_session, original.shift_id, gross=gross, discount=discount, refunds=refund.total_amount
    )
    assert result.net_revenue == summary.net
    assert result.sales_revenue == gross - discount
    assert result.refunds == refund.total_amount
    known = [row for row in result.rows if row.cost_status == "known"]
    assert result.known_gross_profit == sum((row.gross_profit for row in known), Decimal("0"))
    expected = await db_session.scalar(
        text(f"SELECT sum(cost_amount) FROM pos_{kind}_return_slice WHERE command_id=:id"),
        {"id": refund.operation_id},
    )
    assert sum((row.cost for row in known if row.event_kind == "refund"), Decimal("0")) == -expected
    if kind == "physical":
        assert result.cost_status == "known"
        assert result.gross_profit == result.net_revenue - result.net_sales_cost
        if returned == 31:
            assert result.net_revenue == result.net_sales_cost == result.gross_profit == 0
    else:
        assert result.cost_status == "partial" and result.gross_profit is None
        assert len([row for row in result.rows if row.cost_status == "unknown"]) == 2
    assert result.writeoff_cost == 0
    assert await unchanged_facts(db_session) == before
    for offset in range(result.total + 1):
        page = await snapshot(db_session, limit=1, offset=offset)
        assert page.rows == result.rows[offset : offset + 1]
        assert page.net_revenue == result.net_revenue
        assert page.known_gross_profit == result.known_gross_profit
    earlier = await snapshot(db_session, end=refund.receipt.completed_at)
    later = await snapshot(db_session, start=refund.receipt.completed_at)
    assert earlier.refunds == 0 and later.sales_revenue == 0
    assert earlier.net_revenue + later.net_revenue == result.net_revenue


async def test_zero_return_restores_cost_without_changing_cash_reports(
    db_session: AsyncSession,
) -> None:
    source = await mixed_zero(db_session)
    before = await snapshot(db_session)
    await intake(db_session, source)
    after = await snapshot(db_session)
    zero = [row for row in after.rows if row.event_kind == "zero_return"]
    assert len(zero) == 2 and all(row.revenue == 0 for row in zero)
    physical = next(row for row in zero if row.cost_status == "known")
    assert physical.cost < 0 and physical.gross_profit == -physical.cost
    assert after.net_revenue == before.net_revenue and after.refunds == before.refunds == 0
    assert after.known_net_sales_cost == before.known_net_sales_cost + physical.cost
    summary = await reconciled_reports(
        db_session,
        source.original.shift_id,
        gross=checkout_gross(source.original),
        discount=sum((row.discount_amount for row in source.original.lines), Decimal("0")),
        refunds=Decimal("0"),
    )
    assert summary.net == after.net_revenue


@pytest.mark.parametrize("price", [Decimal("0.00"), Decimal("10.00")])
async def test_warehouse_losses_are_separate_and_cancellations_have_no_value(
    db_session: AsyncSession, price: Decimal
) -> None:
    source = await priced_source(db_session, quantity=2, price=price)
    cancelled = uuid4()
    await cancel(db_session, source, cancelled)
    assert (await snapshot(db_session)).total == 0
    await post(db_session, source, selection(closed=1))
    context = runtime(db_session)
    batch = await legacy_batch(db_session, quantity=Decimal("1"))
    legacy = await InventoryService(InventoryRepository(db_session)).write_off(
        batch_id=batch,
        operation_id=uuid4(),
        qty=Decimal("1"),
        reason="damaged",
        comment=None,
        actor_id=context.actor_id,
        tenant_id=context.stock.tenant_id,
        allowed_branch_ids={context.stock.branch_id},
    )
    result = await snapshot(db_session)
    assert result.total == 2 and {row.source_kind for row in result.rows} == {"legacy", "physical"}
    assert all(row.event_kind == "writeoff" and row.gross_profit is None for row in result.rows)
    assert all(row.event_id != cancelled for row in result.rows)
    assert result.writeoff_cost == sum((row.cost for row in result.rows), Decimal("0"))
    assert next(row.cost for row in result.rows if row.source_kind == "legacy") == legacy.amount
    assert result.net_revenue == result.net_sales_cost == result.gross_profit == 0


@pytest.mark.parametrize("visible", [False, True])
async def test_api_export_and_database_permissions_mask_every_cost_field(
    db_session: AsyncSession,
    client: AsyncClient,
    db_engine: AsyncEngine,
    maintenance_engine: AsyncEngine,
    visible: bool,
) -> None:
    context = runtime(db_session)
    permissions = ("reports.view", "reports.export") + (("batches.view_costs",) if visible else ())
    reader = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=permissions
    )
    await physical_cash(db_session, returned=1)
    day = str(await report_day(db_session))
    params = {"date_from": day, "date_to": day, "branch_id": str(context.stock.branch_id)}
    async with report_actor(db_session, reader) as actor:
        response = await client.get(PATH, params={**params, "page_size": "1"})
        assert response.status_code == 200, response.text
        assert response.headers["cache-control"] == "private, no-store"
        data = response.json()
        assert data["total"] == 2 and len(data["rows"]) == 1
        assert data["cost_status"] == ("known" if visible else "hidden")
        for field in (
            "net_sales_cost",
            "known_net_sales_cost",
            "gross_profit",
            "known_gross_profit",
        ):
            assert (data[field] is not None) == visible
        assert (data["rows"][0]["cost"] is not None) == visible
        assert (data["rows"][0]["gross_profit"] is not None) == visible
        exported = await client.get(f"{PATH}.xlsx", params=params)
        assert exported.status_code == 200, exported.text
        sheet = load_workbook(BytesIO(exported.content)).active
        assert sheet is not None and sheet.max_row == 18
        assert (sheet.cell(10, 2).value is not None) == visible
        assert all((sheet.cell(row, 6).value is not None) == visible for row in (17, 18))
        denied = await client.get(PATH, params={**params, "branch_id": str(uuid4())})
        assert denied.status_code == 403
        actor.permissions.remove("reports.export")
        actor.permission_scopes.pop("reports.export")
        assert (await client.get(f"{PATH}.xlsx", params=params)).status_code == 403


async def test_forged_scope_missing_permission_and_readonly_tenant(
    db_session: AsyncSession, db_engine: AsyncEngine, maintenance_engine: AsyncEngine
) -> None:
    context = runtime(db_session)
    reader = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=("customer_returns.view",)
    )
    empty = await snapshot(db_session)
    assert empty.total == 0 and empty.gross_profit == empty.writeoff_cost == 0
    await physical_cash(db_session, returned=1)
    before = await snapshot(db_session)
    assert (await snapshot(db_session, tenant_id=uuid4())).total == 0
    assert (await snapshot(db_session, branch_id=uuid4())).total == 0
    await set_receiving_context(db_session, context.stock, reader)
    assert (await snapshot(db_session)).total == 0
    await db_session.commit()
    async with maintenance_engine.begin() as owner:
        await owner.execute(
            text("UPDATE tenant SET status='readonly' WHERE id=:id"),
            {"id": context.stock.tenant_id},
        )
    await set_receiving_context(db_session, context.stock, context.actor_id)
    assert await snapshot(db_session) == before


async def test_physical_zero_return_has_known_cost_and_no_refund_revenue(
    db_session: AsyncSession,
) -> None:
    await confirmed(db_session, intake_kind="zero_return")
    result = await snapshot(db_session)
    zero = next(row for row in result.rows if row.event_kind == "zero_return")
    assert zero.revenue == 0 and zero.cost_status == "known"
    assert zero.cost < 0 and zero.gross_profit == -zero.cost
    assert result.refunds == 0 and result.cost_status == "known"
    assert result.gross_profit == result.net_revenue - result.net_sales_cost


async def test_training_sale_and_its_refund_are_excluded(db_session: AsyncSession) -> None:
    await physical_cash(db_session, is_test=True)
    result = await snapshot(db_session)
    assert result.total == 0 and result.net_revenue == result.gross_profit == 0
