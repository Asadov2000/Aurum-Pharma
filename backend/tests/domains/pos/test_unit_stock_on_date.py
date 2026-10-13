"""Accepted physical stock shares reports with legacy packs without mixing their units."""

from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from app.domains.foundation.repository import FoundationRepository
from app.domains.foundation.service import FoundationService
from app.domains.inventory.repository import InventoryRepository
from app.domains.inventory.unit_stock_v2 import UnitPackagingV2
from app.domains.pos.repository import POSRepository
from app.domains.pos.schemas import (
    PhysicalStockQuantities,
    StockOnDateData,
    StockOnDateOverview,
    StockRow,
    TopProductPhysicalQuantity,
)
from app.domains.pos.service import POSService
from app.domains.pos.stock_on_date_xlsx import render_stock_on_date_xlsx
from tests.domains.incoming.test_unit_acceptance import db_session as receiving_session_fixture
from tests.domains.incoming.test_unit_acceptance import seed_receiving_actor, set_receiving_context
from tests.domains.inventory.test_unit_cost_ledger_v2 import priced_source
from tests.domains.inventory.test_unit_ledger_v2_dispensing import cart_line, dispense
from tests.domains.inventory.test_unit_physical_blocking import set_blocked
from tests.domains.inventory.test_unit_report import prepare_report_lot
from tests.domains.pos.test_physical_cash_checkout import complete
from tests.domains.pos.test_physical_cash_refund import refund
from tests.domains.pos.test_unit_physical_cart import put_physical
from tests.domains.pos.test_unit_physical_checkout_preparation import checkout_source, saved_request
from tests.domains.pos.test_unit_return import runtime
from tests.domains.pos.test_unit_sale_report_discounts import report_actor, report_day

db_session = receiving_session_fixture


async def overview(
    session: AsyncSession,
    *,
    query: str | None = None,
    page: int = 1,
    page_size: int = 20,
    expires_within_days: int | None = None,
    day: date | None = None,
) -> StockOnDateOverview:
    context = runtime(session)
    return await POSService(POSRepository(session)).get_stock_on_date_overview(
        tenant_id=context.stock.tenant_id,
        on_date=day or await report_day(session),
        branch_id=context.stock.branch_id,
        query=query,
        expires_within_days=expires_within_days,
        page=page,
        page_size=page_size,
    )


async def legacy_batch(
    session: AsyncSession,
    *,
    number: str = "LEGACY MATCH",
    quantity: Decimal = Decimal("0.005"),
) -> UUID:
    context = runtime(session)
    repository = InventoryRepository(session)
    batch = await repository.create_batch(
        tenant_id=context.stock.tenant_id,
        branch_id=context.stock.branch_id,
        catalog_id=context.stock.catalog_id,
        batch_number=number,
        expires_at=date.today() + timedelta(days=10),
        purchase_price=Decimal("3"),
        sale_price=Decimal("10"),
        qty_initial=quantity,
        qty_remaining=Decimal("0"),
    )
    await repository.insert_movement(
        tenant_id=context.stock.tenant_id,
        batch_id=batch.id,
        movement_type="incoming",
        qty_delta=quantity,
    )
    await session.refresh(batch)
    assert batch.qty_remaining == quantity
    return batch.id


async def facts(session: AsyncSession) -> dict[str, list[object]]:
    return {
        table: list(
            (
                await session.execute(
                    text(
                        f"SELECT to_jsonb(row) FROM public.{table} row WHERE tenant_id=:tenant "
                        "ORDER BY to_jsonb(row)::text"
                    ),
                    {"tenant": runtime(session).stock.tenant_id},
                )
            ).scalars()
        )
        for table in (
            "batch",
            "batch_movement",
            "unit_stock_lot_v2",
            "unit_stock_entry_v2",
            "unit_stock_movement_v2",
            "unit_stock_opened_pack_v2",
            "incoming_document",
            "incoming_item",
            "sale",
            "sale_payment",
        )
    }


async def test_accepted_v2_balance_and_original_cost_coexist_with_fractional_legacy(
    db_session: AsyncSession,
) -> None:
    source = await priced_source(db_session, quantity=3, price=Decimal("10"))
    await dispense(
        db_session,
        source,
        (
            cart_line(source, "unit", 5),
            cart_line(source, "blister", 1),
            cart_line(source, "pack", 1),
        ),
    )
    batch_id = await legacy_batch(db_session)
    before = await facts(db_session)
    report = await overview(db_session)
    rows = {row.stock_id: row for row in report.rows}
    physical, legacy = rows[source.lot_id], rows[batch_id]
    assert physical.inventory_kind == "physical" and physical.qty == Decimal("0")
    assert physical.physical_stock is not None
    assert physical.physical_stock.closed_packs == Decimal("1")
    assert physical.physical_stock.opened_units == Decimal("15")
    assert physical.physical_stock.base_units == Decimal("45")
    assert physical.physical_stock.packaging == source.packaging
    assert physical.cost_status == "known" and physical.value == Decimal("15.00")
    assert legacy.inventory_kind == "legacy" and legacy.physical_stock is None
    assert legacy.qty == Decimal("0.005") and legacy.value == Decimal("0.02")
    assert report.total_qty == Decimal("0.005")
    assert [(part.unit, part.quantity) for part in report.physical_quantities] == [
        ("tablet", Decimal("45"))
    ]
    assert report.total_value == report.known_total_value == Decimal("15.02")
    assert report.cost_status == "known"
    assert await facts(db_session) == before


async def test_catalog_changes_and_blocking_do_not_rewrite_historical_physical_stock(
    db_session: AsyncSession,
) -> None:
    source = await priced_source(db_session, quantity=2)
    await dispense(db_session, source, (cart_line(source, "unit", 7),))
    original = await overview(db_session)
    await db_session.execute(
        text(
            "UPDATE tenant_catalog SET base_price=999,unit_sale_preparation="
            "jsonb_set(jsonb_set(unit_sale_preparation,'{units_per_pack}','100'),"
            "'{unit}','\"capsule\"') WHERE id=:id"
        ),
        {"id": runtime(db_session).stock.catalog_id},
    )
    # Blocking is an availability policy; stock still belongs in the balance report.
    await set_blocked(db_session, source, True)
    current = await overview(db_session)
    assert current.rows[0].physical_stock == original.rows[0].physical_stock
    assert current.rows[0].value == original.rows[0].value
    assert current.physical_quantities == original.physical_quantities
    expired = await overview(db_session, day=await report_day(db_session) + timedelta(days=365))
    assert expired.rows[0].physical_stock == original.rows[0].physical_stock


@pytest.mark.parametrize("visible_costs", [False, True])
async def test_reports_only_http_reader_cost_visibility_matches_xlsx_for_both_stock_kinds(
    client: AsyncClient,
    db_session: AsyncSession,
    db_engine: AsyncEngine,
    maintenance_engine: AsyncEngine,
    visible_costs: bool,
) -> None:
    context = runtime(db_session)
    permissions = ("reports.view", "reports.export") + (
        ("batches.view_costs",) if visible_costs else ()
    )
    reader = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=permissions
    )
    await priced_source(db_session, quantity=2, price=Decimal("0"))
    await legacy_batch(db_session)
    before = await facts(db_session)
    params = {"date": str(await report_day(db_session)), "branch_id": str(context.stock.branch_id)}
    async with report_actor(db_session, reader) as actor:
        if visible_costs:
            actor.permissions.add("batches.view_costs")
            actor.permission_scopes["batches.view_costs"] = frozenset({context.stock.branch_id})
        response = await client.get("/api/v1/reports/stock-on-date", params=params)
        assert response.status_code == 200
        payload = response.json()
        assert payload["total"] == 2
        assert payload["cost_status"] == ("known" if visible_costs else "hidden")
        assert all(row["cost_status"] == payload["cost_status"] for row in payload["rows"])
        assert all(
            not {"source_item_id", "catalog_id", "command_id", "lot_id"}.intersection(row)
            for row in payload["rows"]
        )
        if visible_costs:
            assert Decimal(payload["total_value"]) == Decimal("0.02")
            assert Decimal(payload["known_total_value"]) == Decimal("0.02")
        else:
            assert payload["total_value"] is None and payload["known_total_value"] is None
            assert all(
                row["purchase_price"] is None and row["value"] is None for row in payload["rows"]
            )
        exported = await client.get("/api/v1/reports/stock-on-date.xlsx", params=params)
        assert exported.status_code == 200
        sheet = load_workbook(BytesIO(exported.content)).active
        assert sheet is not None
        headers = {cell.value: cell.column for cell in sheet[3]}
        values = [sheet.cell(row, headers["Стоимость"]).value for row in (4, 5)]
        if visible_costs:
            assert sorted(values) == [0, 0.02]
        else:
            assert values == ["Скрыто", "Скрыто"]
        physical_row = next(
            row for row in (4, 5) if sheet.cell(row, headers["Кол-во"]).value == "—"
        )
        assert sheet.cell(physical_row, headers["Базовых единиц"]).value == "60"
    assert await facts(db_session) == before


async def test_filters_precede_paging_and_export_has_all_matching_rows(
    db_session: AsyncSession,
) -> None:
    await priced_source(db_session)
    await legacy_batch(db_session, number="MATCH A")
    await legacy_batch(db_session, number="MATCH B")
    page = await overview(db_session, query="MATCH", page_size=1)
    second = await overview(db_session, query="MATCH", page_size=1, page=2)
    empty = await overview(db_session, query="MATCH", page_size=1, page=3)
    assert page.total == second.total == empty.total == 2
    assert page.total_qty == second.total_qty == empty.total_qty == Decimal("0.010")
    assert page.total_value == second.total_value == empty.total_value == Decimal("0.04")
    assert page.physical_quantities == second.physical_quantities == []
    assert page.rows[0].stock_id != second.rows[0].stock_id and empty.rows == []
    assert (await overview(db_session, query="missing")).total == 0
    assert (await overview(db_session, query="MATCH", expires_within_days=5)).total == 0
    service = POSService(POSRepository(db_session))
    content = await service.get_stock_on_date_xlsx(
        tenant_id=runtime(db_session).stock.tenant_id,
        on_date=await report_day(db_session),
        branch_id=runtime(db_session).stock.branch_id,
        query="MATCH",
        expires_within_days=None,
    )
    sheet = load_workbook(BytesIO(content)).active
    assert sheet is not None
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert [sheet.cell(row, headers["Партия"]).value for row in (4, 5)] == [
        row.batch_number for row in page.rows + second.rows
    ]
    assert sheet.cell(6, 1).value == "ИТОГО"


async def test_refunded_goods_stay_in_quarantine_and_do_not_increase_stock(
    db_session: AsyncSession,
) -> None:
    source = await checkout_source(db_session)
    item = await put_physical(db_session, source.sale.id, sale_format="pack", quantity=1)
    posted = await complete(db_session, await saved_request(db_session, source.sale))
    before = await overview(db_session)
    await refund(db_session, posted.sale_id, item.id, posted.shift_id, quantity=10)
    after = await overview(db_session)
    assert after == before


async def test_date_before_acceptance_has_no_stock_and_valid_empty_workbook(
    db_session: AsyncSession,
) -> None:
    await priced_source(db_session)
    day = await report_day(db_session) - timedelta(days=1)
    report = await overview(db_session, day=day)
    assert report.rows == [] and report.total == 0
    assert report.total_qty == Decimal("0") and report.physical_quantities == []
    assert report.total_value == report.known_total_value == Decimal("0")
    assert report.cost_status == "known"
    content = await POSService(POSRepository(db_session)).get_stock_on_date_xlsx(
        tenant_id=runtime(db_session).stock.tenant_id,
        on_date=day,
        branch_id=runtime(db_session).stock.branch_id,
    )
    assert content[:2] == b"PK"


def test_unknown_stock_value_in_workbook_remains_unknown_with_visible_known_subtotal() -> None:
    packaging = UnitPackagingV2(
        unit="tablet", units_per_pack=30, units_per_blister=10, dispense_step=1
    )
    data = StockOnDateData(
        on_date=date(2026, 10, 13),
        branch_name="Main",
        show_branch_column=False,
        currency="TJS",
        total_qty=Decimal("0.005"),
        total_value=None,
        known_total_value=Decimal("0.02"),
        cost_status="partial",
        physical_quantities=[TopProductPhysicalQuantity(unit="tablet", quantity=Decimal("23"))],
        rows=[
            StockRow(
                name="Historical physical",
                inn=None,
                branch_name="Main",
                batch_number="OLD",
                expires_at=None,
                qty=Decimal("0"),
                inventory_kind="physical",
                physical_stock=PhysicalStockQuantities(
                    packaging=packaging, closed_packs=0, opened_units=23, base_units=23
                ),
                purchase_price=Decimal("10"),
                value=None,
                cost_status="unknown",
            ),
            StockRow(
                name="Known legacy",
                inn=None,
                branch_name="Main",
                batch_number="KNOWN",
                expires_at=None,
                qty=Decimal("0.005"),
                purchase_price=Decimal("3"),
                value=Decimal("0.02"),
            ),
        ],
    )
    sheet = load_workbook(BytesIO(render_stock_on_date_xlsx(data))).active
    assert sheet is not None
    headers = {cell.value: cell.column for cell in sheet[3]}
    assert sheet.cell(4, headers["Стоимость"]).value == "Нет данных"
    assert sheet.cell(4, headers["Кол-во"]).value == "—"
    assert sheet.cell(4, headers["Базовых единиц"]).value == "23"
    totals = next(row for row in range(4, sheet.max_row + 1) if sheet.cell(row, 1).value == "ИТОГО")
    known = next(
        row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, 1).value == "ИЗВЕСТНАЯ СТОИМОСТЬ"
    )
    assert sheet.cell(totals, headers["Стоимость"]).value == "Нет полной стоимости"
    assert sheet.cell(known, headers["Стоимость"]).value == 0.02


async def test_standalone_prepared_v1_stock_is_not_an_accepted_physical_balance(
    db_session: AsyncSession,
) -> None:
    await prepare_report_lot(db_session, cost=Decimal("20"))
    assert (await overview(db_session)).rows == []


async def test_cost_permission_is_evaluated_separately_for_each_visible_branch(
    db_session: AsyncSession, db_engine: AsyncEngine, maintenance_engine: AsyncEngine
) -> None:
    context = runtime(db_session)
    # Every actor and assignment is committed before any receiving transaction
    # takes the tenant access lock; the old cross-connection fixture deadlock cannot recur.
    async with AsyncSession(db_engine) as support:
        async with support.begin():
            await support.execute(text("SELECT set_config('app.support_session','true',true)"))
            branch = await FoundationService(FoundationRepository(support)).create_branch(
                tenant_id=context.stock.tenant_id, fields={"name": "Other stock branch"}
            )
            branch_id = branch.id
    other_stock = replace(context.stock, branch_id=branch_id)
    reader = await seed_receiving_actor(
        db_engine,
        maintenance_engine,
        context.stock,
        permissions=("reports.view", "batches.view_costs"),
    )
    other_reader = await seed_receiving_actor(
        db_engine, maintenance_engine, other_stock, permissions=("reports.view",)
    )
    writer = await seed_receiving_actor(db_engine, maintenance_engine, other_stock)
    async with AsyncSession(db_engine) as support:
        async with support.begin():
            await support.execute(text("SELECT set_config('app.support_session','true',true)"))
            await support.execute(
                text("""
                    INSERT INTO user_assignment
                      (user_id,tenant_id,membership_id,branch_id,role_id,is_active)
                    SELECT :reader,:tenant,own.membership_id,:branch,other.role_id,true
                    FROM user_assignment own CROSS JOIN user_assignment other
                    WHERE own.user_id=:reader AND own.branch_id=:main
                      AND other.user_id=:other AND other.branch_id=:branch
                """),
                {
                    "reader": reader,
                    "tenant": context.stock.tenant_id,
                    "branch": branch_id,
                    "main": context.stock.branch_id,
                    "other": other_reader,
                },
            )
    main = await priced_source(db_session, quantity=2)
    try:
        db_session.info["receiving_runtime"] = replace(context, stock=other_stock, actor_id=writer)
        await set_receiving_context(db_session, other_stock, writer)
        other = await priced_source(db_session, quantity=2)
        db_session.info["receiving_runtime"] = context
        await set_receiving_context(db_session, context.stock, reader)
        report = await POSService(POSRepository(db_session)).build_stock_on_date(
            tenant_id=context.stock.tenant_id, on_date=await report_day(db_session), branch_id=None
        )
        rows = {row.stock_id: row for row in report.rows}
        assert rows[main.lot_id].cost_status == "known" and rows[main.lot_id].value == Decimal("20")
        assert rows[other.lot_id].cost_status == "hidden" and rows[other.lot_id].value is None
        assert rows[other.lot_id].purchase_price is None
        assert report.cost_status == "partial" and report.total_value is None
        assert report.known_total_value == Decimal("20")
    finally:
        db_session.info["receiving_runtime"] = context
        await set_receiving_context(db_session, context.stock, context.actor_id)
