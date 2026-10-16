"""Quarantine valuation is historical, scoped, cost-aware and independent of stock."""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID, uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from app.domains.customer_returns.physical_disposition_repository import (
    PhysicalDispositionRepository,
)
from app.domains.pos.quarantine_report_repository import QuarantineReportRepository
from app.domains.pos.quarantine_report_schemas import QuarantineReportSnapshot
from tests.domains.customer_returns.test_physical_disposition_helpers import (
    confirmed,
    request,
    unchanged_facts,
)
from tests.domains.incoming.test_unit_acceptance import db_session as receiving_session_fixture
from tests.domains.incoming.test_unit_acceptance import seed_receiving_actor, set_receiving_context
from tests.domains.pos.test_unit_return import runtime
from tests.domains.pos.test_unit_sale_report_discounts import report_actor, report_day
from tests.isolation.test_mixed_disposition_migration import confirmed as mixed_confirmed

db_session = receiving_session_fixture
PATH = "/api/v1/reports/quarantine-on-date"


async def snapshot(
    session: AsyncSession,
    *,
    before: datetime | None = None,
    tenant_id: UUID | None = None,
    branch_id: UUID | None = None,
    limit: int | None = None,
    offset: int = 0,
) -> QuarantineReportSnapshot:
    context = runtime(session)
    cutoff = before or await session.scalar(text("SELECT clock_timestamp()+interval '1 day'"))
    assert isinstance(cutoff, datetime)
    return await QuarantineReportRepository(session).snapshot(
        tenant_id=tenant_id or context.stock.tenant_id,
        before=cutoff,
        branch_id=branch_id or context.stock.branch_id,
        limit=limit,
        offset=offset,
    )


@pytest.mark.parametrize("kind", ["physical", "mixed"])
@pytest.mark.parametrize("intake_kind", ["cash_refund", "zero_return"])
async def test_confirmed_intakes_use_saved_slice_cost_and_never_mutate_stock(
    db_session: AsyncSession,
    kind: str,
    intake_kind: Literal["cash_refund", "zero_return"],
) -> None:
    if kind == "physical":
        operation_id = (await confirmed(db_session, intake_kind=intake_kind))[0].operation_id
    else:
        operation_id = (await mixed_confirmed(db_session, intake_kind=intake_kind))[0].operation_id
    before = await unchanged_facts(db_session)
    result = await snapshot(db_session)
    physical = [row for row in result.rows if row.inventory_kind == "physical"]
    assert len(physical) == 1
    row = physical[0]
    assert row.source_kind == kind and row.cost_status == "known"
    assert row.qty == 0 and row.physical_quantity is not None
    expected = await db_session.scalar(
        text(
            f"SELECT sum(cost_amount) FROM pos_{kind}_return_slice "
            "WHERE tenant_id=:tenant AND command_id=:command"
        ),
        {"tenant": runtime(db_session).stock.tenant_id, "command": operation_id},
    )
    assert row.value == expected == result.known_total_value
    if kind == "mixed":
        legacy = [row for row in result.rows if row.inventory_kind == "legacy"]
        assert len(legacy) == 1 and legacy[0].value is None
        assert legacy[0].cost_status == "unknown" and legacy[0].qty > 0
        assert result.cost_status == "partial" and result.total_value is None
        assert result.total_qty == legacy[0].qty
    else:
        assert result.cost_status == "known" and result.total_value == expected
        assert result.total_qty == 0
    assert await unchanged_facts(db_session) == before


async def test_cutoff_uses_intake_and_disposition_history_not_current_queue(
    db_session: AsyncSession,
) -> None:
    source = (await confirmed(db_session))[0]
    assert (await snapshot(db_session, before=source.received_at)).total == 0
    before_decision = await db_session.scalar(text("SELECT clock_timestamp()"))
    assert isinstance(before_decision, datetime)
    pending = await snapshot(db_session, before=before_decision)
    assert pending.total == 1
    resolved = await PhysicalDispositionRepository(db_session).record(
        tenant_id=runtime(db_session).stock.tenant_id, request=request(source)
    )
    assert (await snapshot(db_session, before=before_decision)) == pending
    assert (await snapshot(db_session, before=resolved.resolved_at)).total == 1
    assert (
        await snapshot(db_session, before=resolved.resolved_at + timedelta(microseconds=1))
    ).total == 0


async def test_pagination_keeps_global_valuation_and_separate_quantity_totals(
    db_session: AsyncSession,
) -> None:
    await mixed_confirmed(db_session)
    complete = await snapshot(db_session)
    assert complete.total == 2
    for offset in (0, 1, 2):
        page = await snapshot(db_session, limit=1, offset=offset)
        assert page.total == complete.total
        assert page.total_qty == complete.total_qty
        assert page.physical_quantities == complete.physical_quantities
        assert page.known_total_value == complete.known_total_value
        assert page.total_value is None
        assert page.rows == complete.rows[offset : offset + 1]


@pytest.mark.parametrize("visible", [False, True])
async def test_reports_only_reader_cost_scope_matches_api_and_export(
    db_session: AsyncSession,
    client: AsyncClient,
    db_engine: AsyncEngine,
    maintenance_engine: AsyncEngine,
    visible: bool,
) -> None:
    context = runtime(db_session)
    permissions = ("reports.view", "reports.export") + (("batches.view_costs",) if visible else ())
    reader = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=permissions
    )
    # Assignment guards lock the tenant: seed before the stock transaction holds FK locks.
    await confirmed(db_session)
    params = {"date": str(await report_day(db_session)), "branch_id": str(context.stock.branch_id)}
    async with report_actor(db_session, reader):
        response = await client.get(PATH, params=params)
        assert response.status_code == 200, response.text
        assert response.headers["cache-control"] == "private, no-store"
        data = response.json()
        assert data["total"] == 1
        assert data["cost_status"] == ("known" if visible else "hidden")
        assert (data["rows"][0]["value"] is not None) == visible
        assert (data["known_total_value"] is not None) == visible
        exported = await client.get(f"{PATH}.xlsx", params=params)
        assert exported.status_code == 200, exported.text
        assert exported.headers["cache-control"] == "private, no-store"
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        assert sheet is not None
        assert sheet.cell(4, 2).value == (float(Decimal(data["total_value"])) if visible else None)
        assert sheet.cell(sheet.max_row, 6).value == ("Известна" if visible else "Нет доступа")
        denied = await client.get(PATH, params={**params, "branch_id": str(uuid4())})
        assert denied.status_code == 403


async def test_database_scoping_denies_forged_tenant_branch_and_missing_reports_permission(
    db_session: AsyncSession, db_engine: AsyncEngine, maintenance_engine: AsyncEngine
) -> None:
    context = runtime(db_session)
    reader = await seed_receiving_actor(
        db_engine, maintenance_engine, context.stock, permissions=("customer_returns.view",)
    )
    await confirmed(db_session)
    assert (await snapshot(db_session, tenant_id=uuid4())).total == 0
    assert (await snapshot(db_session, branch_id=uuid4())).total == 0
    await set_receiving_context(db_session, context.stock, reader)
    assert (await snapshot(db_session)).total == 0
    await set_receiving_context(db_session, context.stock, context.actor_id)


async def test_empty_and_readonly_report_are_available_without_mutating_state(
    db_session: AsyncSession, maintenance_engine: AsyncEngine
) -> None:
    empty = await snapshot(db_session)
    assert empty.total == 0 and empty.rows == []
    assert empty.cost_status == "known" and empty.total_value == 0
    await confirmed(db_session)
    before = await snapshot(db_session)
    context = runtime(db_session)
    await db_session.commit()
    async with maintenance_engine.begin() as owner:
        await owner.execute(
            text("UPDATE tenant SET status='readonly' WHERE id=:id"),
            {"id": context.stock.tenant_id},
        )
    await set_receiving_context(db_session, context.stock, context.actor_id)
    assert await snapshot(db_session) == before
