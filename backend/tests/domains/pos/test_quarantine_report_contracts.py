"""Report validation and workbook output never turn unavailable costs into zero."""

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.domains.pos.quarantine_report_schemas import (
    QuarantineReportOverview,
    QuarantineReportRow,
)
from app.domains.pos.quarantine_report_xlsx import render_quarantine_report_xlsx
from app.domains.pos.schemas import TopProductPhysicalQuantity


def row() -> QuarantineReportRow:
    return QuarantineReportRow(
        inventory_kind="physical",
        source_kind="physical",
        intake_operation_id=UUID(int=1),
        line_id=UUID(int=2),
        stock_id=UUID(int=3),
        branch_id=UUID(int=4),
        branch_name="=unsafe",
        name="+unsafe",
        received_at=datetime(2026, 10, 16, tzinfo=UTC),
        qty=Decimal("0"),
        physical_quantity=TopProductPhysicalQuantity(unit="tablet", quantity=Decimal("1")),
        value=Decimal("0.00"),
        cost_status="known",
    )


@pytest.mark.parametrize(
    "patch",
    [
        {"value": None},
        {"cost_status": "hidden"},
        {"cost_status": "unknown"},
        {"qty": "1"},
        {"physical_quantity": None},
        {"value": "-1"},
    ],
)
def test_row_rejects_inconsistent_quantity_or_cost_visibility(patch: dict[str, object]) -> None:
    with pytest.raises(ValidationError):
        QuarantineReportRow.model_validate(row().model_dump() | patch)


def test_known_zero_is_preserved_and_workbook_escapes_untrusted_names() -> None:
    source = row()
    assert source.physical_quantity is not None
    overview = QuarantineReportOverview(
        on_date=date(2026, 10, 16),
        rows=[source],
        total=1,
        page=1,
        page_size=25,
        total_qty=Decimal("0"),
        physical_quantities=[source.physical_quantity],
        total_value=Decimal("0.00"),
        known_total_value=Decimal("0.00"),
        cost_status="known",
    )
    sheet = load_workbook(BytesIO(render_quarantine_report_xlsx(overview))).active
    assert sheet is not None
    assert sheet.cell(4, 2).value == 0 and sheet.cell(4, 2).number_format == "0.00"
    assert sheet.cell(sheet.max_row, 1).value == "'+unsafe"
    assert sheet.cell(sheet.max_row, 2).value == "'=unsafe"
    assert sheet.cell(sheet.max_row, 5).value == 0
