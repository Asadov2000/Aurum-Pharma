"""Stock-on-date: ledger reconstruction (later movements excluded), valuation,
and the empty → valid-file case."""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.domains.inventory.repository import InventoryRepository
from app.domains.pos.repository import POSRepository
from app.domains.pos.service import POSService
from tests.platform_access_helpers import create_test_platform_user


async def _stock_reader(session: AsyncSession, tenant_id: UUID) -> None:
    reader = await create_test_platform_user(session, access_kind="developer")
    await session.execute(
        text(
            "SELECT set_config('app.user_id',:actor,true),"
            "set_config('app.tenant_id',:tenant,true)"
        ),
        {"actor": str(reader.id), "tenant": str(tenant_id)},
    )


async def test_fractional_stock_totals_and_filtered_export_match_screen(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold(batch_qty=Decimal("0.005"))
    await _stock_reader(db_session, s["tenant"].id)
    inventory = InventoryRepository(db_session)
    today = await _today(db_session)
    second = await inventory.create_batch(
        tenant_id=s["tenant"].id,
        branch_id=s["branch"].id,
        catalog_id=s["item"].id,
        batch_number="MATCH-ROUNDING",
        expires_at=today + timedelta(days=10),
        purchase_price=Decimal("3.00"),
        sale_price=Decimal("10.00"),
        qty_initial=Decimal("0.005"),
        qty_remaining=Decimal("0"),
    )
    await inventory.insert_movement(
        tenant_id=s["tenant"].id,
        batch_id=second.id,
        movement_type="incoming",
        qty_delta=Decimal("0.005"),
        source_table=None,
        source_id=None,
    )
    service = POSService(POSRepository(db_session))
    data = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=today, branch_id=None
    )
    overview = await service.get_stock_on_date_overview(
        tenant_id=s["tenant"].id,
        on_date=today,
        branch_id=None,
        query=None,
        expires_within_days=None,
        page=1,
        page_size=50,
    )
    assert data.total_value == overview.total_value == Decimal("0.04")
    assert all(row.value is not None for row in overview.rows)
    assert sum(row.value for row in overview.rows if row.value is not None) == overview.total_value
    s["branch"].is_active = False
    await db_session.flush()
    archived = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=today, branch_id=None
    )
    assert archived.show_branch_column is True
    assert all(row.branch_name == s["branch"].name for row in archived.rows)
    for query, expiry, expected_rows in [
        ("MATCH", 30, 1),
        ("missing", 30, 0),
        ("MATCH", 5, 0),
    ]:
        overview = await service.get_stock_on_date_overview(
            tenant_id=s["tenant"].id,
            on_date=today,
            branch_id=None,
            query=query,
            expires_within_days=expiry,
            page=1,
            page_size=50,
        )
        content = await service.get_stock_on_date_xlsx(
            tenant_id=s["tenant"].id,
            on_date=today,
            branch_id=None,
            query=query,
            expires_within_days=expiry,
        )
        sheet = load_workbook(BytesIO(content)).active
        assert sheet is not None
        assert overview.total == expected_rows
        assert sheet.max_row == 4 + expected_rows
        value_column = next(cell.column for cell in sheet[3] if cell.value == "Стоимость")
        assert overview.total_value is not None
        assert sheet.cell(sheet.max_row, value_column).value == float(overview.total_value)


async def _today(db: AsyncSession) -> date:
    # Reports interpret dates in the tenant tz (Asia/Dushanbe by default), so
    # "today" here must be the local date — not UTC, which disagrees by a day
    # between 19:00–23:59 UTC.
    return (
        await db.execute(text("SELECT (now() AT TIME ZONE 'Asia/Dushanbe')::date"))
    ).scalar_one()


async def _sell(service: POSService, s, qty: Decimal, paid: Decimal) -> None:  # type: ignore[no-untyped-def]
    sale = await service.create_sale(
        tenant_id=s["tenant"].id, register_id=s["register"].id, cashier_user_id=s["cashier"].id
    )
    await service.add_item(sale_id=sale.id, catalog_id=s["item"].id, qty=qty)
    await service.add_payment(sale_id=sale.id, payment_method="cash", amount=paid)
    await service.complete(sale_id=sale.id)


async def test_stock_on_date_value_and_totals(db_session: AsyncSession, pos_scaffold) -> None:
    # batch: +100 in, purchase_price 3.00. Sell 30 → 70 remain.
    s = await pos_scaffold(batch_qty=100, sale_price=10)
    await _stock_reader(db_session, s["tenant"].id)
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    await _sell(service, s, qty=Decimal("30"), paid=Decimal("300"))

    data = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=await _today(db_session), branch_id=None
    )
    assert len(data.rows) == 1
    row = data.rows[0]
    assert row.qty == Decimal("70.000")
    assert row.purchase_price == Decimal("3.00")
    assert row.value == Decimal("210.00")  # 70 × 3.00
    assert data.total_qty == Decimal("70.000")
    assert data.total_value == Decimal("210.00")


async def test_stock_on_date_excludes_movements_after_the_date(
    db_session: AsyncSession, pos_scaffold
) -> None:
    s = await pos_scaffold(
        batch_qty=100,
        sale_price=10,
        movement_created_at=datetime.now(UTC) - timedelta(days=5),
    )
    await _stock_reader(db_session, s["tenant"].id)
    service = POSService(POSRepository(db_session))
    await service.open_shift(
        tenant_id=s["tenant"].id,
        register_id=s["register"].id,
        opened_by_user_id=s["cashier"].id,
        opening_cash=Decimal("0"),
    )
    await _sell(service, s, qty=Decimal("40"), paid=Decimal("400"))

    today = await _today(db_session)
    # 3 days ago: only the incoming counts → full 100 (the later sale is excluded).
    historical = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=today - timedelta(days=3), branch_id=None
    )
    assert len(historical.rows) == 1
    assert historical.rows[0].qty == Decimal("100.000")
    assert historical.total_value == Decimal("300.00")  # 100 × 3.00

    # Today: incoming + sale → 60.
    current = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=today, branch_id=None
    )
    assert current.rows[0].qty == Decimal("60.000")
    assert current.total_value == Decimal("180.00")


async def test_stock_on_date_empty_is_valid_file(db_session: AsyncSession, pos_scaffold) -> None:
    s = await pos_scaffold(batch_qty=100)
    await _stock_reader(db_session, s["tenant"].id)
    service = POSService(POSRepository(db_session))
    today = await _today(db_session)

    # 10 days before the only (incoming) movement → nothing on hand.
    data = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=today - timedelta(days=10), branch_id=None
    )
    assert data.rows == []
    assert data.total_qty == Decimal("0")
    assert data.total_value == Decimal("0")

    xlsx = await service.get_stock_on_date_xlsx(
        tenant_id=s["tenant"].id, on_date=today - timedelta(days=10), branch_id=None
    )
    assert xlsx[:2] == b"PK"
    assert len(xlsx) > 500


async def test_stock_on_date_uses_tenant_timezone(db_session: AsyncSession, pos_scaffold) -> None:
    """A movement at 01:00 Dushanbe (= 20:00 UTC the previous day) counts toward
    the local day's stock, not the previous UTC day. Default tz Asia/Dushanbe."""
    s = await pos_scaffold(
        batch_qty=100,
        movement_created_at=datetime(2026, 6, 14, 20, 0, tzinfo=UTC),
    )
    await _stock_reader(db_session, s["tenant"].id)
    service = POSService(POSRepository(db_session))

    # Local day 2026-06-15 sees the incoming → stock present.
    on_15 = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=date(2026, 6, 15), branch_id=None
    )
    assert len(on_15.rows) == 1
    assert on_15.rows[0].qty == Decimal("100.000")

    # The previous UTC day (06-14) must NOT see it (it's local 06-15).
    on_14 = await service.build_stock_on_date(
        tenant_id=s["tenant"].id, on_date=date(2026, 6, 14), branch_id=None
    )
    assert on_14.rows == []
